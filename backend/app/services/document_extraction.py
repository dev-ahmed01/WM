"""Untrusted document *preview* extraction. Never publishes or executes imported text.

An arbitrary document is not a validated operational workflow. The publishing API
continues to require a reviewed OWD document and its deterministic validation gate.
"""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_PREVIEW_CHARACTERS = 12000
MAX_ARCHIVE_EXPANDED_BYTES = 120 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 3000
MAX_PDF_PAGES = 200
MAX_WORKBOOK_ROWS = 2000
MAX_WORKBOOK_CELLS = 20000

SUPPORTED_DOCUMENT_EXTENSIONS = frozenset({
    ".md", ".txt", ".csv", ".tsv", ".json", ".html", ".htm",
    ".pdf", ".docx", ".xlsx", ".pptx", ".rtf",
})
OFFICE_EXTENSIONS = frozenset({".docx", ".xlsx", ".pptx"})


class DocumentExtractionError(ValueError):
    """An actionable failure safe to display to administrators."""


@dataclass(frozen=True)
class DocumentPreview:
    filename: str
    file_type: str
    text_preview: str
    total_characters: int
    preview_truncated: bool
    warnings: list[str]


def _text_bytes(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError(
            "This text file is not UTF-8. Export it as UTF-8 and try again."
        ) from exc


def _check_archive(raw: bytes) -> None:
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise DocumentExtractionError("The Office document has too many archive entries.")
            total = sum(member.file_size for member in members)
            if total > MAX_ARCHIVE_EXPANDED_BYTES:
                raise DocumentExtractionError("The Office document expands beyond the safety limit.")
            if any(
                member.file_size > 2 * 1024 * 1024
                and member.compress_size > 0
                and member.file_size / member.compress_size > 200
                for member in members
            ):
                raise DocumentExtractionError("The Office document has an unsafe compression ratio.")
    except BadZipFile as exc:
        raise DocumentExtractionError("This Office file is damaged or not a valid Office document.") from exc


def _extract_text(filename: str, raw: bytes) -> tuple[str, list[str]]:
    extension = Path(filename).suffix.lower()
    warnings: list[str] = []

    if extension in OFFICE_EXTENSIONS:
        _check_archive(raw)

    if extension in {".txt", ".md"}:
        return _text_bytes(raw), warnings

    if extension in {".csv", ".tsv"}:
        separator = "\t" if extension == ".tsv" else ","
        rows = csv.reader(io.StringIO(_text_bytes(raw)), delimiter=separator)
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows), warnings

    if extension == ".json":
        try:
            data = json.loads(_text_bytes(raw))
        except json.JSONDecodeError as exc:
            raise DocumentExtractionError("This JSON document is not valid.") from exc
        return json.dumps(data, ensure_ascii=False, indent=2), warnings

    if extension in {".html", ".htm"}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(_text_bytes(raw), "html.parser")
        for tag in soup(["script", "style", "iframe", "noscript"]):
            tag.decompose()
        return soup.get_text(separator="\n", strip=True), warnings

    if extension == ".rtf":
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(_text_bytes(raw)), warnings

    if extension == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(raw), strict=False)
            if reader.is_encrypted:
                raise DocumentExtractionError("Password-protected PDFs require an unlocked copy.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise DocumentExtractionError(
                    f"PDF exceeds the {MAX_PDF_PAGES}-page preview safety limit."
                )
            content = "\n\n".join(
                f"[Page {index}]\n{page.extract_text() or ''}"
                for index, page in enumerate(reader.pages, start=1)
            )
        except DocumentExtractionError:
            raise
        except Exception as exc:
            raise DocumentExtractionError("This PDF could not be read.") from exc
        if len(content.strip()) < 40:
            warnings.append("Little or no selectable text found; scanned pages need OCR review.")
        return content, warnings

    if extension == ".docx":
        from docx import Document
        try:
            doc = Document(io.BytesIO(raw))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(paragraphs), warnings
        except Exception as exc:
            raise DocumentExtractionError("This Word file could not be read.") from exc

    if extension == ".xlsx":
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts: list[str] = []
            total_cells = 0
            try:
                for sheet in book.worksheets[:20]:
                    parts.append(f"[Sheet: {sheet.title}]")
                    for row in sheet.iter_rows(max_row=MAX_WORKBOOK_ROWS, values_only=True):
                        total_cells += len(row)
                        if total_cells > MAX_WORKBOOK_CELLS:
                            warnings.append("Spreadsheet preview limited to the first 20,000 cells.")
                            break
                        rendered = [str(value) if value is not None else "" for value in row]
                        if any(rendered):
                            parts.append(" | ".join(rendered))
                    if total_cells > MAX_WORKBOOK_CELLS:
                        break
                if len(book.worksheets) > 20:
                    warnings.append("Only the first 20 sheets are previewed.")
            finally:
                book.close()
            return "\n".join(parts), warnings
        except Exception as exc:
            raise DocumentExtractionError("This spreadsheet could not be read.") from exc

    if extension == ".pptx":
        from pptx import Presentation
        try:
            slides = Presentation(io.BytesIO(raw)).slides
            parts = []
            for i, slide in enumerate(slides, start=1):
                parts.append(f"[Slide {i}]")
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text.strip():
                        parts.append(shape.text.strip())
                    if shape.has_table:
                        for row in shape.table.rows:
                            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(parts), warnings
        except Exception as exc:
            raise DocumentExtractionError("This presentation could not be read.") from exc

    raise DocumentExtractionError(
        "Unsupported format. Use PDF, DOCX, XLSX, PPTX, TXT, MD, CSV, TSV, JSON, HTML or RTF."
    )


def preview_document(filename: str, raw: bytes) -> DocumentPreview:
    """Read a bounded document in memory; return only a short administrator preview.

    No external requests, persistent files, vector indexing, graph compilation or
    publication occur here. Preview text must never be treated as approved guidance.
    """
    safe_name = Path(filename.replace("\\", "/")).name
    extension = Path(safe_name).suffix.lower()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise DocumentExtractionError(
            "Unsupported format. Use PDF, DOCX, XLSX, PPTX, TXT, MD, CSV, TSV, JSON, HTML or RTF."
        )
    if not raw:
        raise DocumentExtractionError("The document is empty.")
    if len(raw) > MAX_DOCUMENT_BYTES:
        raise DocumentExtractionError("Document exceeds the 25 MB upload limit.")
    try:
        content, warnings = _extract_text(safe_name, raw)
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError("Document extraction failed. Check the file and try again.") from exc
    content = content.strip()
    if not content:
        raise DocumentExtractionError(
            "No readable text was extracted. Image-only documents require OCR before import."
        )
    return DocumentPreview(
        filename=safe_name,
        file_type=extension.lstrip("."),
        text_preview=content[:MAX_PREVIEW_CHARACTERS],
        total_characters=len(content),
        preview_truncated=len(content) > MAX_PREVIEW_CHARACTERS,
        warnings=warnings,
    )
